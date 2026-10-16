"""CSV (streamed) and XLSX (write-only, constant memory) writers for tabular datasets."""
from __future__ import annotations

import csv
import io
from typing import Iterable, Sequence

from django.http import HttpResponse, StreamingHttpResponse
from django.utils import timezone

XLSX_ROW_LIMIT = 250_000
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe(value):
    """Neutralise spreadsheet formula injection in text cells (CSV/Excel)."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


class _Echo:
    def write(self, value):
        return value


def csv_response(filename: str, headers: Sequence[str], keys: Sequence[str], rows: Iterable[dict], meta: dict) -> StreamingHttpResponse:
    writer = csv.writer(_Echo())

    def stream():
        yield "﻿"  # UTF-8 BOM so Excel detects the encoding
        yield writer.writerow(headers)
        for row in rows:
            yield writer.writerow([_safe(row.get(k)) if row.get(k) is not None else "" for k in keys])

    response = StreamingHttpResponse(stream(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    for key, value in meta.items():
        response[f"X-Report-{key}"] = str(value)
    return response


def xlsx_response(filename: str, title: str, headers: Sequence[str], keys: Sequence[str], rows: Iterable[dict],
                  meta: dict) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    bold = Font(bold=True)
    header_cells = []
    for h in headers:
        cell = WriteOnlyCell(ws, value=h)
        cell.font = bold
        header_cells.append(cell)
    ws.append(header_cells)
    count = 0
    for row in rows:
        count += 1
        if count > XLSX_ROW_LIMIT:
            break
        ws.append([_safe(row.get(k)) for k in keys])
    info = wb.create_sheet("Report info")
    info.append(["Report", title])
    for key, value in meta.items():
        info.append([key, str(value)])
    info.append(["Rows", min(count, XLSX_ROW_LIMIT)])
    if count > XLSX_ROW_LIMIT:
        info.append(["Truncated", f"Only the first {XLSX_ROW_LIMIT} rows are included; use CSV for the full dataset."])
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def stamp() -> str:
    return timezone.now().strftime("%Y%m%d-%H%M")
